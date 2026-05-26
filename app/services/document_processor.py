import logging
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select
from sqlalchemy.orm import selectinload
from app.models.document import Document, DocumentChunk, DocumentProcessingAttempt, ManualRefinedDocument
from app.models.enums import ProcessingStatus, ProcessingCapability, ProtectionType
from app.core.config import settings
from app.core.storage import S3Storage
from typing import Optional, List, Tuple
import hashlib
import httpx
import fitz  # pymupdf

logger = logging.getLogger(__name__)


class DocumentProcessor:
    def __init__(self, db: AsyncSession, storage: Optional[S3Storage] = None):
        self.db = db
        self.storage = storage or self._get_storage()
        self.embedding_api_url = settings.EMBEDDING_API_URL

    def _get_storage(self) -> S3Storage:
        from app.core.storage import storage as default_storage
        return default_storage

    async def process(self, document_id: int):
        logger.info(f"Processing document {document_id}")

        result = await self.db.execute(
            select(Document).where(Document.id == document_id)
        )
        document = result.scalar_one_or_none()

        if not document:
            logger.error(f"Document {document_id} not found")
            return

        attempt = DocumentProcessingAttempt(
            document_id=document_id,
            attempt_type="full_processing"
        )
        self.db.add(attempt)

        try:
            document.processing_status = ProcessingStatus.METADATA_EXTRACTED
            document.protection_type = ProtectionType.NONE

            # Check if manual refined content exists
            has_refined = await self._check_refined(document)
            if has_refined:
                document.is_refined = True
                document.processing_capability = ProcessingCapability.FULLTEXT_EXTRACTABLE
                document.processing_status = ProcessingStatus.REFINED_UPLOADED
                logger.info(f"Document {document_id} has refined content")
                await self.db.commit()
                return

            # Extract actual text content from file
            content = await self._extract_content(document)

            if not content or not content.strip():
                document.processing_status = ProcessingStatus.AWAITING_MANUAL_REFINEMENT
                document.processing_capability = ProcessingCapability.MANUAL_REFINED_ONLY
                logger.warning(f"Document {document_id} requires manual refinement (no text extracted)")
                await self.db.commit()
                return

            # Chunk the content
            chunks = self._chunk_content(content)
            logger.info(f"Document {document_id}: {len(chunks)} chunks created")

            # Generate embeddings and store chunks
            await self._create_chunks_with_embeddings(document, chunks)

            document.processing_capability = ProcessingCapability.EMBEDDABLE
            document.processing_status = ProcessingStatus.READY_FOR_SEARCH
            attempt.status = "success"
            logger.info(f"Document {document_id} processed successfully ({len(chunks)} chunks)")

        except Exception as e:
            logger.error(f"Error processing document {document_id}: {e}", exc_info=True)
            document.processing_status = ProcessingStatus.PROCESSING_FAILED
            document.processing_error_reason = str(e)
            attempt.status = "failed"
            attempt.error_message = str(e)

        await self.db.commit()

    async def _check_refined(self, document: Document) -> bool:
        result = await self.db.execute(
            select(ManualRefinedDocument)
            .where(ManualRefinedDocument.original_document_id == document.id)
            .limit(1)
        )
        return result.scalar_one_or_none() is not None

    async def _extract_content(self, document: Document) -> Optional[str]:
        """Download file from S3 and extract text content."""
        if not document.file_path:
            logger.warning(f"Document {document.id} has no file_path")
            return None

        try:
            file_bytes = await self.storage.download_file(document.file_path)
        except Exception as e:
            logger.error(f"Failed to download file from S3: {e}")
            return None

        mime = document.mime_type or ""
        filename = (document.original_filename or document.file_path or "").lower()

        try:
            # PDF extraction using pymupdf
            if mime == "application/pdf" or filename.endswith(".pdf"):
                return self._extract_pdf(file_bytes)

            # Plain text
            if mime.startswith("text/") or filename.endswith((".txt", ".md", ".csv")):
                return self._extract_text(file_bytes)

            # Word documents
            if "wordprocessingml" in mime or filename.endswith((".docx", ".doc")):
                return self._extract_docx(file_bytes)

            # Excel spreadsheets
            if "spreadsheetml" in mime or "ms-excel" in mime or filename.endswith((".xlsx", ".xls")):
                return self._extract_xlsx(file_bytes, filename)

            # PowerPoint presentations
            if "presentationml" in mime or "ms-powerpoint" in mime or filename.endswith((".pptx", ".ppt")):
                return self._extract_pptx(file_bytes)

            # Fallback: try as text
            logger.warning(f"Unknown mime type '{mime}' for document {document.id}, trying text decode")
            return self._extract_text(file_bytes)

        except Exception as e:
            logger.error(f"Content extraction failed for document {document.id}: {e}")
            return None

    def _extract_pdf(self, file_bytes: bytes) -> str:
        """Extract text from PDF using pymupdf (fitz)."""
        text_parts = []
        with fitz.open(stream=file_bytes, filetype="pdf") as pdf_doc:
            for page_num, page in enumerate(pdf_doc):
                page_text = page.get_text("text")
                if page_text.strip():
                    text_parts.append(page_text)

        full_text = "\n\n".join(text_parts)
        logger.info(f"PDF extraction: {len(text_parts)} pages, {len(full_text)} chars")
        return full_text

    def _extract_text(self, file_bytes: bytes) -> str:
        """Decode plain text files."""
        for encoding in ("utf-8", "euc-kr", "cp949", "latin-1"):
            try:
                return file_bytes.decode(encoding)
            except (UnicodeDecodeError, LookupError):
                continue
        return file_bytes.decode("utf-8", errors="replace")

    def _extract_docx(self, file_bytes: bytes) -> str:
        """Extract text from .docx files."""
        import io
        from docx import Document as DocxDocument

        doc = DocxDocument(io.BytesIO(file_bytes))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        return "\n\n".join(paragraphs)

    def _extract_xlsx(self, file_bytes: bytes, filename: str = "") -> str:
        """Extract text from .xlsx/.xls files.

        Each sheet is converted to readable text with rows represented
        as 'header: value' pairs so the content is embedding-friendly.
        """
        import io

        # .xls (legacy) via xlrd
        if filename.endswith(".xls"):
            import xlrd
            wb = xlrd.open_workbook(file_contents=file_bytes)
            parts: list[str] = []
            for sheet in wb.sheets():
                if sheet.nrows == 0:
                    continue
                header = [str(sheet.cell_value(0, c)).strip() for c in range(sheet.ncols)]
                sheet_lines = [f"[시트: {sheet.name}]"]
                for r in range(1, sheet.nrows):
                    row_vals = [str(sheet.cell_value(r, c)).strip() for c in range(sheet.ncols)]
                    if not any(row_vals):
                        continue
                    pairs = [f"{h}: {v}" for h, v in zip(header, row_vals) if v]
                    if pairs:
                        sheet_lines.append(" | ".join(pairs))
                if len(sheet_lines) > 1:
                    parts.append("\n".join(sheet_lines))
            text = "\n\n".join(parts)
            logger.info(f"XLS extraction: {wb.nsheets} sheets, {len(text)} chars")
            return text

        # .xlsx via openpyxl
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
        parts: list[str] = []

        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            rows = list(ws.iter_rows(values_only=True))
            if not rows:
                continue

            # Find header row: first row with 2+ non-empty cells
            header_idx = 0
            for i, row in enumerate(rows):
                non_empty = sum(1 for c in row if c is not None and str(c).strip())
                if non_empty >= 2:
                    header_idx = i
                    break

            header = [str(c).strip() if c is not None else "" for c in rows[header_idx]]
            sheet_lines = [f"[시트: {sheet_name}]"]

            for row in rows[header_idx + 1:]:
                vals = [str(c).strip() if c is not None else "" for c in row]
                if not any(vals):
                    continue
                # Skip summary/total rows
                first_val = vals[0].lower() if vals[0] else ""
                if any(kw in first_val for kw in ("합계", "소계", "total", "sum", "계")):
                    continue
                pairs = [f"{h}: {v}" for h, v in zip(header, vals) if v and h]
                if pairs:
                    sheet_lines.append(" | ".join(pairs))

            if len(sheet_lines) > 1:
                parts.append("\n".join(sheet_lines))

        wb.close()
        text = "\n\n".join(parts)
        logger.info(f"XLSX extraction: {len(wb.sheetnames)} sheets, {len(text)} chars")
        return text

    def _extract_pptx(self, file_bytes: bytes) -> str:
        """Extract text from .pptx files (slide text + speaker notes)."""
        import io
        from pptx import Presentation

        prs = Presentation(io.BytesIO(file_bytes))
        parts: list[str] = []

        for slide_num, slide in enumerate(prs.slides, 1):
            slide_texts: list[str] = []

            # Shape text (titles, text boxes, tables)
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for para in shape.text_frame.paragraphs:
                        text = para.text.strip()
                        if text:
                            slide_texts.append(text)
                if shape.has_table:
                    table = shape.table
                    for row in table.rows:
                        row_text = " | ".join(
                            cell.text.strip() for cell in row.cells if cell.text.strip()
                        )
                        if row_text:
                            slide_texts.append(row_text)

            # Speaker notes
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                notes = slide.notes_slide.notes_text_frame.text.strip()
                if notes:
                    slide_texts.append(f"(노트: {notes})")

            if slide_texts:
                parts.append(f"[슬라이드 {slide_num}]\n" + "\n".join(slide_texts))

        text = "\n\n".join(parts)
        logger.info(f"PPTX extraction: {len(prs.slides)} slides, {len(text)} chars")
        return text

    def _chunk_content(self, content: str) -> List[str]:
        """Split content into overlapping chunks."""
        chunks = []
        chunk_size = settings.CHUNK_SIZE
        overlap = settings.CHUNK_OVERLAP

        # Clean up whitespace
        content = content.strip()
        if not content:
            return []

        for i in range(0, len(content), chunk_size - overlap):
            chunk = content[i:i + chunk_size]
            if chunk.strip():
                chunks.append(chunk.strip())

        return chunks or [content]

    async def _create_chunks_with_embeddings(
        self, document: Document, chunks: List[str]
    ):
        """Create document chunks with dense and sparse embeddings from TEI."""
        batch_size = 32

        for batch_start in range(0, len(chunks), batch_size):
            batch = chunks[batch_start:batch_start + batch_size]
            dense_vectors, sparse_vectors = await self._get_embeddings(batch)

            for idx, chunk_text in enumerate(batch):
                global_idx = batch_start + idx
                chunk = DocumentChunk(
                    document_id=document.id,
                    customer_id=document.customer_id,
                    document_title=document.title,
                    content=chunk_text,
                    chunk_index=global_idx,
                    token_count=len(chunk_text.split()),
                    metadata_json='{"source": "auto_extracted"}',
                )

                # Set embeddings if available
                if dense_vectors and idx < len(dense_vectors):
                    chunk.dense_vector = dense_vectors[idx]
                if sparse_vectors and idx < len(sparse_vectors):
                    chunk.sparse_vector = sparse_vectors[idx]

                self.db.add(chunk)

    async def _get_embeddings(
        self, texts: List[str]
    ) -> Tuple[List[List[float]], List[Optional[str]]]:
        """Get dense and sparse embeddings from TEI (bge-m3) service.
        
        Returns:
            dense: List of dense vectors
            sparse: List of SPARSEVEC-formatted strings for pgvector
        """
        dense = []
        sparse: List[Optional[str]] = []
        sparse_dim = 250002  # BGE-m3 vocab size

        try:
            async with httpx.AsyncClient(timeout=120.0) as client:
                # Dense embeddings
                response = await client.post(
                    f"{self.embedding_api_url}/embed",
                    json={"inputs": texts, "truncate": True}
                )
                response.raise_for_status()
                data = response.json()
                dense = data if isinstance(data, list) else data.get("embeddings", [])

                # Sparse embeddings (bge-m3 supports this)
                try:
                    sparse_response = await client.post(
                        f"{self.embedding_api_url}/embed_sparse",
                        json={"inputs": texts, "truncate": True}
                    )
                    if sparse_response.status_code == 200:
                        raw_sparse = sparse_response.json()
                        for sv in raw_sparse:
                            if not sv:
                                sparse.append(None)
                                continue
                            parts = []
                            for entry in sv:
                                idx = entry.get("index", 0)
                                val = entry.get("value", 0.0)
                                if val != 0.0:
                                    parts.append(f"{idx}:{val}")
                            sparse.append("{" + ",".join(parts) + "}/" + str(sparse_dim))
                except Exception:
                    logger.debug("Sparse embedding endpoint not available, skipping")

        except Exception as e:
            logger.error(f"Embedding generation failed: {e}")

        return dense, sparse
